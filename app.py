# -*- coding: utf-8 -*-
"""
العليان لإدارة المبيعات
يدعم SQLite محلياً و PostgreSQL على Vercel
"""
import os
import io
import secrets
from datetime import datetime, timedelta, date
from functools import wraps

from flask import (Flask, render_template, request, jsonify, send_file, g)
from werkzeug.security import generate_password_hash, check_password_hash
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

# ============ الإعدادات ============
app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", secrets.token_hex(32))

# اكتشاف البيئة: PostgreSQL على Vercel، SQLite محلياً
DATABASE_URL = (
    os.environ.get("POSTGRES_URL")
    or os.environ.get("DATABASE_URL")
    or os.environ.get("STORAGE_URL")
    or os.environ.get("DATABASE_URL_UNPOOLED")
    or os.environ.get("POSTGRES_URL_NON_POOLING")
)
USE_POSTGRES = bool(DATABASE_URL)

if USE_POSTGRES:
    import psycopg
    from psycopg.rows import dict_row
else:
    import sqlite3

DB_PATH = os.path.join(os.path.dirname(__file__), "database.db")


# ============ قاعدة البيانات ============
class DBWrapper:
    """يوحّد واجهة SQLite و PostgreSQL"""
    def __init__(self, conn, is_pg):
        self.conn = conn
        self.is_pg = is_pg

    def execute(self, query, params=None):
        params = params or []
        if self.is_pg:
            query = query.replace("?", "%s")
            query = query.replace("INSERT OR IGNORE", "INSERT")
            # ON CONFLICT للـ products
            if "INSERT INTO products" in query and "ON CONFLICT" not in query:
                query += " ON CONFLICT (name) DO NOTHING"
            cur = self.conn.cursor(row_factory=dict_row)
            cur.execute(query, params)
            return cur
        else:
            cur = self.conn.cursor()
            cur.execute(query, params)
            return cur

    def commit(self):
        self.conn.commit()


def get_db():
    if "db" not in g:
        if USE_POSTGRES:
            conn = psycopg.connect(DATABASE_URL)
            g.db = DBWrapper(conn, True)
        else:
            conn = sqlite3.connect(DB_PATH)
            conn.row_factory = sqlite3.Row
            g.db = DBWrapper(conn, False)
    return g.db


@app.teardown_appcontext
def close_db(exc):
    db = g.pop("db", None)
    if db is not None:
        db.conn.close()


def init_db():
    """إنشاء الجداول - يعمل مع SQLite و PostgreSQL"""
    if USE_POSTGRES:
        conn = psycopg.connect(DATABASE_URL, autocommit=True)
        cur = conn.cursor()
    else:
        conn = sqlite3.connect(DB_PATH)
        cur = conn.cursor()

    # تحديد نوع العمود الذاتي حسب البيئة
    PK = "SERIAL PRIMARY KEY" if USE_POSTGRES else "INTEGER PRIMARY KEY AUTOINCREMENT"
    TS = "TIMESTAMP DEFAULT CURRENT_TIMESTAMP"

    cur.execute(f"""
        CREATE TABLE IF NOT EXISTS users (
            id {PK},
            username TEXT UNIQUE NOT NULL,
            password TEXT NOT NULL,
            role TEXT NOT NULL,
            full_name TEXT,
            active INTEGER DEFAULT 1,
            created_at {TS}
        )
    """)

    cur.execute(f"""
        CREATE TABLE IF NOT EXISTS products (
            id {PK},
            name TEXT UNIQUE NOT NULL,
            last_price REAL DEFAULT 0
        )
    """)

    cur.execute(f"""
        CREATE TABLE IF NOT EXISTS invoices (
            id {PK},
            type TEXT NOT NULL,
            invoice_no TEXT,
            party TEXT,
            date TEXT NOT NULL,
            total REAL DEFAULT 0,
            notes TEXT,
            user_id INTEGER,
            created_at {TS}
        )
    """)

    cur.execute(f"""
        CREATE TABLE IF NOT EXISTS invoice_items (
            id {PK},
            invoice_id INTEGER NOT NULL,
            product_name TEXT NOT NULL,
            quantity REAL NOT NULL,
            price REAL NOT NULL,
            subtotal REAL NOT NULL
        )
    """)

    cur.execute(f"""
        CREATE TABLE IF NOT EXISTS expenses (
            id {PK},
            category TEXT NOT NULL,
            amount REAL NOT NULL,
            date TEXT NOT NULL,
            notes TEXT,
            user_id INTEGER,
            created_at {TS}
        )
    """)

    cur.execute(f"""
        CREATE TABLE IF NOT EXISTS tokens (
            token TEXT PRIMARY KEY,
            user_id INTEGER NOT NULL,
            created_at {TS}
        )
    """)

    cur.execute(f"""
        CREATE TABLE IF NOT EXISTS activity_log (
            id {PK},
            user_id INTEGER,
            action TEXT,
            details TEXT,
            created_at {TS}
        )
    """)

    # مستخدمون افتراضيون
    cur.execute("SELECT COUNT(*) FROM users")
    count = cur.fetchone()[0]
    if count == 0:
        cur.execute(
            "INSERT INTO users (username,password,role,full_name) VALUES (?,?,?,?)" if not USE_POSTGRES
            else "INSERT INTO users (username,password,role,full_name) VALUES (%s,%s,%s,%s)",
            ("admin", generate_password_hash("admin123"), "manager", "المدير العام"))
        cur.execute(
            "INSERT INTO users (username,password,role,full_name) VALUES (?,?,?,?)" if not USE_POSTGRES
            else "INSERT INTO users (username,password,role,full_name) VALUES (%s,%s,%s,%s)",
            ("emp", generate_password_hash("emp123"), "employee", "موظف المبيعات"))
        print("✅ مستخدمون افتراضيون: admin/admin123 - emp/emp123")

    if not USE_POSTGRES:
        conn.commit()
    conn.close()


# ============ المصادقة ============
def get_user_by_token(token):
    if not token:
        return None
    db = get_db()
    row = db.execute("""
        SELECT u.* FROM users u
        JOIN tokens t ON t.user_id = u.id
        WHERE t.token = ? AND u.active = 1
    """, (token,)).fetchone()
    return dict(row) if row else None


def token_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        token = request.args.get("t", "").strip()
        user = get_user_by_token(token)
        if not user:
            return jsonify({"ok": False, "error": "غير مصرح"}), 401
        request.user = user
        request.token = token
        return f(*args, **kwargs)
    return wrapper


def manager_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if getattr(request, "user", {}).get("role") != "manager":
            return jsonify({"ok": False, "error": "هذه الصفحة للمدير فقط"}), 403
        return f(*args, **kwargs)
    return wrapper


def log_action(user_id, action, details=""):
    db = get_db()
    db.execute("INSERT INTO activity_log (user_id,action,details) VALUES (?,?,?)",
               (user_id, action, details))
    db.commit()


# ============ أدوات مساعدة ============
def calc_range(period, custom_from=None, custom_to=None):
    today = date.today()
    if period == "day":
        return today.isoformat(), today.isoformat()
    if period == "week":
        start = today - timedelta(days=today.weekday())
        return start.isoformat(), today.isoformat()
    if period == "month":
        return today.replace(day=1).isoformat(), today.isoformat()
    if period == "year":
        return today.replace(month=1, day=1).isoformat(), today.isoformat()
    if period == "custom" and custom_from and custom_to:
        return custom_from, custom_to
    return "1900-01-01", "2999-12-31"


def compute_totals(date_from, date_to, user_filter=None):
    db = get_db()
    uc_inv, uc_exp = "", ""
    up_inv, up_exp = [], []
    if user_filter:
        uc_inv = " AND user_id = ?"
        uc_exp = " AND user_id = ?"
        up_inv = [user_filter]
        up_exp = [user_filter]

    sales = db.execute(
        f"SELECT COALESCE(SUM(total),0) AS v FROM invoices WHERE type='sale' AND date BETWEEN ? AND ?{uc_inv}",
        [date_from, date_to] + up_inv).fetchone()["v"]
    purchases = db.execute(
        f"SELECT COALESCE(SUM(total),0) AS v FROM invoices WHERE type='purchase' AND date BETWEEN ? AND ?{uc_inv}",
        [date_from, date_to] + up_inv).fetchone()["v"]
    expenses = db.execute(
        f"SELECT COALESCE(SUM(amount),0) AS v FROM expenses WHERE date BETWEEN ? AND ?{uc_exp}",
        [date_from, date_to] + up_exp).fetchone()["v"]

    sold_q = f"""
        SELECT ii.product_name, SUM(ii.quantity) AS qty
        FROM invoice_items ii
        JOIN invoices i ON i.id = ii.invoice_id
        WHERE i.type='sale' AND i.date BETWEEN ? AND ?{uc_inv.replace('user_id', 'i.user_id')}
        GROUP BY ii.product_name
    """
    sold = db.execute(sold_q, [date_from, date_to] + up_inv).fetchall()

    cogs = 0.0
    missing = []
    for item in sold:
        avg_row = db.execute("""
            SELECT COALESCE(AVG(ii.price), 0) AS v
            FROM invoice_items ii JOIN invoices i ON i.id = ii.invoice_id
            WHERE i.type='purchase' AND ii.product_name = ?
        """, (item["product_name"],)).fetchone()
        avg_cost = float(avg_row["v"]) if avg_row else 0
        if avg_cost == 0:
            missing.append(item["product_name"])
        cogs += avg_cost * float(item["qty"])

    return {
        "sales": float(sales),
        "purchases": float(purchases),
        "expenses": float(expenses),
        "cogs": cogs,
        "profit": float(sales) - cogs - float(expenses),
        "missing_cost": missing,
    }


# ============ الصفحة الرئيسية ============
@app.route("/")
def index():
    # تأكد من وجود الجداول أول مرة
    try:
        init_db()
    except Exception as e:
        print("init_db:", e)
    return render_template("index.html")


# ============ API: الدخول ============
@app.route("/api/login", methods=["POST"])
def api_login():
    data = request.get_json() or {}
    username = data.get("username", "").strip()
    password = data.get("password", "")
    role = data.get("role", "")

    db = get_db()
    row = db.execute("SELECT * FROM users WHERE username=?", (username,)).fetchone()
    if not row:
        return jsonify({"ok": False, "error": "اسم المستخدم أو كلمة المرور خاطئة"}), 401
    row = dict(row)
    if not check_password_hash(row["password"], password):
        return jsonify({"ok": False, "error": "اسم المستخدم أو كلمة المرور خاطئة"}), 401
    if not row["active"]:
        return jsonify({"ok": False, "error": "الحساب معطّل"}), 403
    if row["role"] != role:
        return jsonify({"ok": False, "error": "الدور المختار لا يطابق الحساب"}), 403

    token = secrets.token_urlsafe(24)
    db.execute("INSERT INTO tokens (token, user_id) VALUES (?, ?)", (token, row["id"]))
    db.commit()
    log_action(row["id"], "login", f"{username} دخل")

    return jsonify({
        "ok": True,
        "token": token,
        "user": {
            "id": row["id"],
            "username": row["username"],
            "full_name": row["full_name"],
            "role": row["role"],
        }
    })


@app.route("/api/logout", methods=["POST"])
@token_required
def api_logout():
    db = get_db()
    db.execute("DELETE FROM tokens WHERE token = ?", (request.token,))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/me")
@token_required
def api_me():
    return jsonify({"ok": True, "user": request.user})


# ============ API: لوحة التحكم ============
@app.route("/api/dashboard")
@token_required
@manager_required
def api_dashboard():
    period = request.args.get("period", "month")
    user_filter = request.args.get("user_id")
    user_filter = int(user_filter) if user_filter and user_filter.strip() else None
    date_from, date_to = calc_range(period, request.args.get("from"), request.args.get("to"))

    db = get_db()
    totals = compute_totals(date_from, date_to, user_filter)

    uc = " AND user_id = ?" if user_filter else ""
    up = [user_filter] if user_filter else []

    recent_sales = db.execute(
        f"SELECT id,date,party,total FROM invoices WHERE type='sale'{uc} ORDER BY id DESC LIMIT 5",
        up).fetchall()
    recent_purchases = db.execute(
        f"SELECT id,date,party,total FROM invoices WHERE type='purchase'{uc} ORDER BY id DESC LIMIT 5",
        up).fetchall()

    labels, s_data, p_data = [], [], []
    for i in range(29, -1, -1):
        d = (date.today() - timedelta(days=i)).isoformat()
        s = db.execute(
            f"SELECT COALESCE(SUM(total),0) AS v FROM invoices WHERE type='sale' AND date=?{uc}",
            [d] + up).fetchone()["v"]
        p = db.execute(
            f"SELECT COALESCE(SUM(total),0) AS v FROM invoices WHERE type='purchase' AND date=?{uc}",
            [d] + up).fetchone()["v"]
        labels.append(d[5:])
        s_data.append(float(s))
        p_data.append(float(p))

    users = db.execute("SELECT id,username,full_name,role FROM users ORDER BY id").fetchall()

    activity = db.execute("""
        SELECT u.id, u.username, u.full_name, u.role,
            (SELECT COALESCE(SUM(total),0) FROM invoices WHERE type='sale' AND user_id=u.id AND date BETWEEN ? AND ?) AS sales_total,
            (SELECT COALESCE(SUM(total),0) FROM invoices WHERE type='purchase' AND user_id=u.id AND date BETWEEN ? AND ?) AS purchase_total,
            (SELECT COALESCE(SUM(amount),0) FROM expenses WHERE user_id=u.id AND date BETWEEN ? AND ?) AS expense_total
        FROM users u ORDER BY u.role DESC, u.id
    """, (date_from, date_to, date_from, date_to, date_from, date_to)).fetchall()

    return jsonify({
        "ok": True,
        "totals": totals,
        "date_from": date_from,
        "date_to": date_to,
        "recent_sales": [dict(r) for r in recent_sales],
        "recent_purchases": [dict(r) for r in recent_purchases],
        "chart_labels": labels,
        "chart_sales": s_data,
        "chart_purchases": p_data,
        "users": [dict(u) for u in users],
        "activity": [dict(a) for a in activity],
    })


# ============ API: المنتجات ============
@app.route("/api/products")
@token_required
def api_products():
    q = request.args.get("q", "").strip()
    db = get_db()
    if q:
        rows = db.execute("SELECT name, last_price FROM products WHERE name LIKE ? LIMIT 8",
                          (f"%{q}%",)).fetchall()
    else:
        rows = db.execute("SELECT name, last_price FROM products ORDER BY id DESC LIMIT 8").fetchall()
    return jsonify([{"name": r["name"], "price": r["last_price"]} for r in rows])


# ============ API: حفظ فاتورة ============
@app.route("/api/invoice/<itype>", methods=["POST"])
@token_required
def api_invoice(itype):
    if itype not in ("sale", "purchase"):
        return jsonify({"ok": False, "error": "نوع غير صحيح"}), 400

    data = request.get_json() or {}
    party = data.get("party", "").strip()
    notes = data.get("notes", "").strip()
    inv_no = data.get("invoice_no", "").strip()
    items = data.get("items", [])

    valid = []
    for it in items:
        try:
            n = str(it.get("product_name", "")).strip()
            q = float(it.get("quantity", 0))
            p = float(it.get("price", 0))
        except (ValueError, TypeError):
            continue
        if n and q > 0 and p >= 0:
            valid.append((n, q, p, q * p))

    if not valid:
        return jsonify({"ok": False, "error": "أضف منتجاً واحداً على الأقل"}), 400

    total = sum(v[3] for v in valid)
    today = date.today().isoformat()
    db = get_db()

    cur = db.execute(
        "INSERT INTO invoices (type,invoice_no,party,date,total,notes,user_id) VALUES (?,?,?,?,?,?,?) RETURNING id"
        if USE_POSTGRES else
        "INSERT INTO invoices (type,invoice_no,party,date,total,notes,user_id) VALUES (?,?,?,?,?,?,?)",
        (itype, inv_no, party, today, total, notes, request.user["id"]))

    if USE_POSTGRES:
        inv_id = cur.fetchone()["id"]
    else:
        inv_id = cur.lastrowid

    for n, q, p, sub in valid:
        db.execute("INSERT INTO invoice_items (invoice_id,product_name,quantity,price,subtotal) VALUES (?,?,?,?,?)",
                   (inv_id, n, q, p, sub))
        # upsert للمنتج
        existing = db.execute("SELECT id FROM products WHERE name=?", (n,)).fetchone()
        if existing:
            db.execute("UPDATE products SET last_price=? WHERE name=?", (p, n))
        else:
            db.execute("INSERT INTO products (name,last_price) VALUES (?,?)", (n, p))

    db.commit()
    log_action(request.user["id"], f"add_{itype}", f"فاتورة #{inv_id} بمبلغ {total:.2f}")
    return jsonify({"ok": True, "invoice_id": inv_id, "total": total})


# ============ API: مصروف ============
@app.route("/api/expense", methods=["POST"])
@token_required
def api_expense():
    data = request.get_json() or {}
    category = data.get("category", "").strip()
    try:
        amount = float(data.get("amount", 0))
    except (ValueError, TypeError):
        amount = 0
    notes = data.get("notes", "").strip()

    if not category or amount <= 0:
        return jsonify({"ok": False, "error": "أدخل تصنيفاً ومبلغاً صحيحاً"}), 400

    db = get_db()
    db.execute("INSERT INTO expenses (category,amount,date,notes,user_id) VALUES (?,?,?,?,?)",
               (category, amount, date.today().isoformat(), notes, request.user["id"]))
    db.commit()
    log_action(request.user["id"], "add_expense", f"{category} بمبلغ {amount:.2f}")
    return jsonify({"ok": True})


# ============ API: العمليات ============
@app.route("/api/invoices/<itype>")
@token_required
def api_invoices_list(itype):
    if itype not in ("sale", "purchase"):
        return jsonify({"ok": False, "error": "نوع غير صحيح"}), 400

    limit = int(request.args.get("limit", 10))
    db = get_db()

    if request.user["role"] == "manager":
        rows = db.execute(
            "SELECT i.*, u.username, u.full_name FROM invoices i "
            "LEFT JOIN users u ON u.id = i.user_id "
            "WHERE i.type=? ORDER BY i.id DESC LIMIT ?", (itype, limit)).fetchall()
    else:
        rows = db.execute(
            "SELECT i.*, u.username, u.full_name FROM invoices i "
            "LEFT JOIN users u ON u.id = i.user_id "
            "WHERE i.type=? AND i.user_id=? ORDER BY i.id DESC LIMIT ?",
            (itype, request.user["id"], limit)).fetchall()

    invoices = []
    for inv in rows:
        items = db.execute("SELECT * FROM invoice_items WHERE invoice_id=?", (inv["id"],)).fetchall()
        d = dict(inv)
        d["items"] = [dict(i) for i in items]
        invoices.append(d)

    return jsonify({"ok": True, "invoices": invoices})


@app.route("/api/expenses")
@token_required
def api_expenses_list():
    limit = int(request.args.get("limit", 10))
    db = get_db()
    if request.user["role"] == "manager":
        rows = db.execute(
            "SELECT e.*, u.username, u.full_name FROM expenses e "
            "LEFT JOIN users u ON u.id = e.user_id "
            "ORDER BY e.id DESC LIMIT ?", (limit,)).fetchall()
    else:
        rows = db.execute(
            "SELECT e.*, u.username, u.full_name FROM expenses e "
            "LEFT JOIN users u ON u.id = e.user_id "
            "WHERE e.user_id=? ORDER BY e.id DESC LIMIT ?",
            (request.user["id"], limit)).fetchall()
    return jsonify({"ok": True, "expenses": [dict(r) for r in rows]})


# ============ API: التقارير ============
@app.route("/api/reports")
@token_required
@manager_required
def api_reports():
    period = request.args.get("period", "month")
    user_filter = request.args.get("user_id")
    user_filter = int(user_filter) if user_filter and user_filter.strip() else None
    date_from, date_to = calc_range(period, request.args.get("from"), request.args.get("to"))

    db = get_db()
    totals = compute_totals(date_from, date_to, user_filter)

    uc = " AND user_id = ?" if user_filter else ""
    up = [user_filter] if user_filter else []

    sales = db.execute(
        f"SELECT id,date,party,total FROM invoices WHERE type='sale' AND date BETWEEN ? AND ?{uc} ORDER BY id DESC",
        [date_from, date_to] + up).fetchall()
    purchases = db.execute(
        f"SELECT id,date,party,total FROM invoices WHERE type='purchase' AND date BETWEEN ? AND ?{uc} ORDER BY id DESC",
        [date_from, date_to] + up).fetchall()
    exps = db.execute(
        f"SELECT id,category,amount,date,notes FROM expenses WHERE date BETWEEN ? AND ?{uc} ORDER BY id DESC",
        [date_from, date_to] + up).fetchall()

    top_products = db.execute(f"""
        SELECT ii.product_name, SUM(ii.quantity) AS qty, SUM(ii.subtotal) AS total
        FROM invoice_items ii JOIN invoices i ON i.id=ii.invoice_id
        WHERE i.type='sale' AND i.date BETWEEN ? AND ?{uc.replace('user_id','i.user_id')}
        GROUP BY ii.product_name ORDER BY total DESC LIMIT 5
    """, [date_from, date_to] + up).fetchall()

    return jsonify({
        "ok": True,
        "totals": totals,
        "date_from": date_from,
        "date_to": date_to,
        "sales": [dict(r) for r in sales],
        "purchases": [dict(r) for r in purchases],
        "expenses": [dict(r) for r in exps],
        "top_products": [dict(r) for r in top_products],
    })


# ============ API: تصدير Excel ============
@app.route("/api/export/excel/<kind>")
@token_required
@manager_required
def api_export_excel(kind):
    period = request.args.get("period", "month")
    user_filter = request.args.get("user_id")
    user_filter = int(user_filter) if user_filter and user_filter.strip() else None
    date_from, date_to = calc_range(period, request.args.get("from"), request.args.get("to"))

    db = get_db()
    uc = " AND i.user_id = ?" if user_filter else ""
    uc_exp = " AND user_id = ?" if user_filter else ""
    up = [user_filter] if user_filter else []

    wb = Workbook()
    ws = wb.active
    ws.sheet_view.rightToLeft = True
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True, size=12)

    if kind in ("sales", "purchases"):
        itype = "sale" if kind == "sales" else "purchase"
        ws.title = "المبيعات" if kind == "sales" else "المشتريات"
        headers = ["التاريخ", "رقم الفاتورة", "الاسم", "المنتج", "الكمية", "السعر", "الإجمالي الفرعي", "المستخدم"]
        ws.append(headers)
        for col in range(1, len(headers) + 1):
            c = ws.cell(row=1, column=col)
            c.fill = header_fill; c.font = header_font
            c.alignment = Alignment(horizontal="center")

        rows = db.execute(f"""
            SELECT i.date, i.invoice_no, i.party, ii.product_name, ii.quantity, ii.price, ii.subtotal,
                   u.full_name, u.username
            FROM invoice_items ii
            JOIN invoices i ON i.id = ii.invoice_id
            LEFT JOIN users u ON u.id = i.user_id
            WHERE i.type=? AND i.date BETWEEN ? AND ?{uc}
            ORDER BY i.id, ii.id
        """, [itype, date_from, date_to] + up).fetchall()

        grand = 0
        for r in rows:
            ws.append([r["date"], r["invoice_no"] or "-", r["party"] or "-",
                       r["product_name"], r["quantity"], f"{r['price']:.2f}",
                       f"{r['subtotal']:.2f}", r["full_name"] or r["username"] or "-"])
            grand += float(r["subtotal"])
        ws.append([])
        ws.append(["", "", "", "", "", "", f"{grand:.2f}", "الإجمالي"])
    elif kind == "expenses":
        ws.title = "المصروفات"
        headers = ["التصنيف", "المبلغ", "التاريخ", "المستخدم", "ملاحظات"]
        ws.append(headers)
        for col in range(1, len(headers) + 1):
            c = ws.cell(row=1, column=col)
            c.fill = header_fill; c.font = header_font
            c.alignment = Alignment(horizontal="center")

        rows = db.execute(f"""
            SELECT e.category, e.amount, e.date, e.notes, u.full_name, u.username
            FROM expenses e LEFT JOIN users u ON u.id = e.user_id
            WHERE e.date BETWEEN ? AND ?{uc_exp}
            ORDER BY e.id
        """, [date_from, date_to] + up).fetchall()

        grand = 0
        for r in rows:
            ws.append([r["category"], f"{float(r['amount']):.2f}", r["date"],
                       r["full_name"] or r["username"] or "-", r["notes"] or "-"])
            grand += float(r["amount"])
        ws.append([])
        ws.append(["", f"{grand:.2f}", "", "", "الإجمالي"])
    else:
        return jsonify({"ok": False, "error": "نوع غير صحيح"}), 400

    for col in ws.columns:
        try:
            ws.column_dimensions[col[0].column_letter].width = 18
        except Exception:
            pass

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return send_file(bio, as_attachment=True,
                     download_name=f"{kind}_{date_from}_{date_to}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ============ API: بيانات الطباعة ============
@app.route("/api/print/<kind>")
@token_required
@manager_required
def api_print(kind):
    period = request.args.get("period", "month")
    user_filter = request.args.get("user_id")
    user_filter = int(user_filter) if user_filter and user_filter.strip() else None
    date_from, date_to = calc_range(period, request.args.get("from"), request.args.get("to"))

    db = get_db()
    uc = " AND i.user_id = ?" if user_filter else ""
    uc_exp = " AND user_id = ?" if user_filter else ""
    up = [user_filter] if user_filter else []

    if kind in ("sales", "purchases"):
        itype = "sale" if kind == "sales" else "purchase"
        title = "تقرير المبيعات" if kind == "sales" else "تقرير المشتريات"

        rows = db.execute(f"""
            SELECT i.id, i.date, i.invoice_no, i.party, i.total,
                   ii.product_name, ii.quantity, ii.price, ii.subtotal,
                   u.full_name, u.username
            FROM invoice_items ii
            JOIN invoices i ON i.id = ii.invoice_id
            LEFT JOIN users u ON u.id = i.user_id
            WHERE i.type=? AND i.date BETWEEN ? AND ?{uc}
            ORDER BY i.id, ii.id
        """, [itype, date_from, date_to] + up).fetchall()

        invoices = {}
        for r in rows:
            iid = r["id"]
            if iid not in invoices:
                invoices[iid] = {
                    "id": iid, "date": r["date"], "invoice_no": r["invoice_no"],
                    "party": r["party"], "total": float(r["total"]),
                    "user": r["full_name"] or r["username"] or "-",
                    "items": []
                }
            invoices[iid]["items"].append({
                "product_name": r["product_name"],
                "quantity": r["quantity"],
                "price": float(r["price"]),
                "subtotal": float(r["subtotal"]),
            })

        grand = sum(inv["total"] for inv in invoices.values())
        return jsonify({"ok": True, "kind": kind, "title": title,
                        "invoices": list(invoices.values()),
                        "total": grand,
                        "date_from": date_from, "date_to": date_to})
    elif kind == "expenses":
        rows = db.execute(f"""
            SELECT e.category, e.amount, e.date, e.notes, u.full_name, u.username
            FROM expenses e LEFT JOIN users u ON u.id = e.user_id
            WHERE e.date BETWEEN ? AND ?{uc_exp}
            ORDER BY e.id
        """, [date_from, date_to] + up).fetchall()
        grand = sum(float(r["amount"]) for r in rows)
        return jsonify({"ok": True, "kind": "expenses", "title": "تقرير المصروفات",
                        "expenses": [dict(r) for r in rows],
                        "total": grand,
                        "date_from": date_from, "date_to": date_to})

    return jsonify({"ok": False, "error": "نوع غير صحيح"}), 400


# ============ API: المستخدمون ============
@app.route("/api/users", methods=["GET", "POST"])
@token_required
@manager_required
def api_users():
    db = get_db()
    if request.method == "POST":
        data = request.get_json() or {}
        action = data.get("action")
        if action == "add":
            username = data.get("username", "").strip()
            password = data.get("password", "")
            role = data.get("role", "employee")
            full_name = data.get("full_name", "").strip()
            if not username or not password:
                return jsonify({"ok": False, "error": "أدخل اسم مستخدم وكلمة مرور"}), 400
            existing = db.execute("SELECT id FROM users WHERE username=?", (username,)).fetchone()
            if existing:
                return jsonify({"ok": False, "error": "اسم المستخدم موجود"}), 400
            db.execute("INSERT INTO users (username,password,role,full_name) VALUES (?,?,?,?)",
                       (username, generate_password_hash(password), role, full_name))
            db.commit()
            log_action(request.user["id"], "add_user", f"إضافة {username}")
            return jsonify({"ok": True})
        elif action == "toggle":
            uid = int(data.get("user_id"))
            row = db.execute("SELECT active FROM users WHERE id=?", (uid,)).fetchone()
            if row:
                db.execute("UPDATE users SET active=? WHERE id=?",
                           (0 if row["active"] else 1, uid))
                db.commit()
                return jsonify({"ok": True})
        elif action == "reset":
            uid = int(data.get("user_id"))
            newpass = data.get("new_password", "")
            if newpass:
                db.execute("UPDATE users SET password=? WHERE id=?",
                           (generate_password_hash(newpass), uid))
                db.commit()
                return jsonify({"ok": True})
        return jsonify({"ok": False, "error": "إجراء غير معروف"}), 400

    users = db.execute("SELECT id,username,full_name,role,active,created_at FROM users ORDER BY id").fetchall()
    logs = db.execute("SELECT * FROM activity_log ORDER BY id DESC LIMIT 20").fetchall()
    return jsonify({"ok": True,
                    "users": [dict(u) for u in users],
                    "logs": [dict(l) for l in logs]})


# ============ تشغيل محلي ============
if __name__ == "__main__":
    init_db()
    app.run(host="0.0.0.0", port=5000, debug=True)